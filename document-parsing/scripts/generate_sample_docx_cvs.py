#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""generate_sample_docx_cvs.py -- produit les CV DOCX de
document-parsing/fixtures/cv_samples/ utilises par one-pager/tests (voie texte
DocIE pour les DOCX, #180).

Trois CV fictifs, MEME contenu, trois mises en forme Word courantes, construits
avec python-docx (deja dependance de cv-parser/requirements.txt) :

  1. cv_docx_titres.docx   -- styles Titre 1/2, puces Word (List Bullet),
                              competences et formation en tableaux.
  2. cv_docx_sans_titres.docx -- aucun style de titre : sections en MAJUSCULES
                              grasses, retours a la ligne manuels (Maj+Entree)
                              et tabulations entre intitule et dates.
  3. cv_docx_tableau.docx  -- modele « barre laterale » : tableau 1x2, contact /
                              competences / langues a gauche, missions a droite
                              dans un tableau imbrique dates | detail.

Usage :
    python generate_sample_docx_cvs.py [output_dir]
    (defaut : document-parsing/fixtures/cv_samples)
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

FIXTURES_DIR = Path(__file__).parent.parent / "fixtures" / "cv_samples"

NOM = "Alice Dupont"
TITRE = "Data Engineer senior"
EMAIL = "alice.dupont@example.com"
TEL = "06 12 34 56 78"
VILLE = "Lille"
MISSIONS = [
    dict(role="Data Engineer", company="Decathlon", dates="Mars 2022 - Aujourd'hui", lieu="Lille",
         bullets=["Conception de la plateforme data sur Databricks pour 40 equipes",
                  "Reduction de 35 % des couts de calcul par optimisation des jobs Spark"],
         env="Python, Databricks, Spark, Azure"),
    dict(role="Developpeuse BI", company="Leroy Merlin", dates="Septembre 2018 - Fevrier 2022", lieu="Lezennes",
         bullets=["Migration de 120 rapports Business Objects vers Power BI",
                  "Animation d'ateliers de recueil de besoins avec les metiers"],
         env="SQL Server, Power BI, SSIS"),
]
COMPETENCES = [("Langages", "Python, SQL, Scala"), ("Cloud", "Azure, AWS"), ("Outils", "Databricks, Airflow, Power BI")]
FORMATION = [("2018", "Master Informatique decisionnelle", "Universite de Lille")]
LANGUES = [("Anglais", "courant (C1)"), ("Espagnol", "notions")]


def _gras(doc, texte, taille=None):
    p = doc.add_paragraph()
    r = p.add_run(texte)
    r.bold = True
    if taille:
        r.font.size = Pt(taille)
    return p


def cv_titres(dest: Path) -> None:
    d = Document()
    d.add_heading(NOM, level=0)
    d.add_paragraph(TITRE)
    d.add_paragraph(f"{EMAIL} | {TEL} | {VILLE}")
    d.add_heading("Experiences professionnelles", level=1)
    for m in MISSIONS:
        d.add_heading(f"{m['role']} - {m['company']}", level=2)
        d.add_paragraph(f"{m['dates']} | {m['lieu']}")
        for b in m["bullets"]:
            d.add_paragraph(b, style="List Bullet")
        d.add_paragraph(f"Environnement technique : {m['env']}")
    d.add_heading("Competences", level=1)
    t = d.add_table(rows=len(COMPETENCES), cols=2)
    for i, (cat, items) in enumerate(COMPETENCES):
        t.cell(i, 0).text = cat
        t.cell(i, 1).text = items
    d.add_heading("Formation", level=1)
    t = d.add_table(rows=len(FORMATION), cols=3)
    for i, ligne in enumerate(FORMATION):
        for j, v in enumerate(ligne):
            t.cell(i, j).text = v
    d.add_heading("Langues", level=1)
    for lang, niv in LANGUES:
        d.add_paragraph(f"{lang} : {niv}", style="List Bullet")
    d.save(dest)


def cv_sans_titres(dest: Path) -> None:
    d = Document()
    _gras(d, NOM, 18)
    d.add_paragraph(TITRE)
    p = d.add_paragraph(EMAIL)
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run(TEL)
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run(VILLE)
    _gras(d, "EXPERIENCES PROFESSIONNELLES")
    for m in MISSIONS:
        p = d.add_paragraph()
        p.add_run(m["role"]).bold = True
        p.add_run("\t" + m["dates"])
        p.add_run().add_break(WD_BREAK.LINE)
        p.add_run(f"{m['company']}, {m['lieu']}")
        for b in m["bullets"]:
            d.add_paragraph("- " + b)
        d.add_paragraph(f"Environnement technique : {m['env']}")
    _gras(d, "COMPETENCES")
    for cat, items in COMPETENCES:
        d.add_paragraph(f"{cat} : {items}")
    _gras(d, "FORMATION")
    for an, dip, eco in FORMATION:
        d.add_paragraph(f"{an}\t{dip}, {eco}")
    _gras(d, "LANGUES")
    for lang, niv in LANGUES:
        d.add_paragraph(f"{lang} : {niv}")
    d.save(dest)


def cv_tableau(dest: Path) -> None:
    d = Document()
    _gras(d, NOM, 18)
    d.add_paragraph(TITRE)
    mise_en_page = d.add_table(rows=1, cols=2)
    gauche, droite = mise_en_page.cell(0, 0), mise_en_page.cell(0, 1)
    gauche.paragraphs[0].text = "CONTACT"
    for v in (EMAIL, TEL, VILLE):
        gauche.add_paragraph(v)
    gauche.add_paragraph("COMPETENCES")
    for cat, items in COMPETENCES:
        gauche.add_paragraph(f"{cat} : {items}")
    gauche.add_paragraph("LANGUES")
    for lang, niv in LANGUES:
        gauche.add_paragraph(f"{lang} : {niv}")
    droite.paragraphs[0].text = "EXPERIENCES"
    missions = droite.add_table(rows=len(MISSIONS), cols=2)
    for i, m in enumerate(MISSIONS):
        missions.cell(i, 0).text = m["dates"]
        c = missions.cell(i, 1)
        c.paragraphs[0].text = f"{m['role']} - {m['company']} ({m['lieu']})"
        for b in m["bullets"]:
            c.add_paragraph("- " + b)
        c.add_paragraph(f"Environnement technique : {m['env']}")
    droite.add_paragraph("FORMATION")
    for an, dip, eco in FORMATION:
        droite.add_paragraph(f"{an} - {dip}, {eco}")
    d.save(dest)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else FIXTURES_DIR
    out.mkdir(parents=True, exist_ok=True)
    cv_titres(out / "cv_docx_titres.docx")
    cv_sans_titres(out / "cv_docx_sans_titres.docx")
    cv_tableau(out / "cv_docx_tableau.docx")
    print(f"3 DOCX ecrits dans {out}")


if __name__ == "__main__":
    main()
