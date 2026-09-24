"""Le texte d'un .docx envoyé à DocIE garde ses séparateurs.

`docie_client.document_payload` rend le texte qu'on poste sur
`/v1/extract/text` ; c'est de ce texte, et de lui seul, que DocIE extrait la
fiche qui entre dans la CVthèque. L'ancien lecteur tenait en une ligne :

    "\\n".join("".join(p.itertext()) for p in root.findall(".//w:p", ns))

`itertext()` ne rend rien pour les éléments vides `w:br` (Maj+Entrée), `w:cr`
et `w:tab` : le texte de part et d'autre était COLLÉ
(`alice.dupont@example.com06 12 34 56 78Lille`). Il rendait aussi le code des
champs (`w:instrText`) et le texte des révisions supprimées (`w:delText`),
séparait chaque cellule de tableau de sa catégorie, et comptait QUATRE fois une
zone de texte (le paragraphe d'ancrage rend `mc:Choice` + `mc:Fallback`, puis
`.//w:p` retrouve les paragraphes intérieurs des deux branches).

Critère de mesure, le même que #188 côté one-pager : une valeur connue
« survit » si elle est présente avec un séparateur non alphanumérique (ou le
bord du texte) de chaque côté.

Construction des entrées, générées ici et jamais versionnées :
- `python-docx` : retours manuels (`run.add_break()`), tabulations
  (`run.add_tab()`, plus un taquet de tabulation dans `w:pPr`), tableaux simples
  et tableau de mise en page avec tableau imbriqué ;
- XML `word/document.xml` écrit À LA MAIN dans un zip (python-docx ne sait pas
  les produire) : `w:cr`, zone de texte `mc:AlternateContent`, codes de champ
  `w:fldChar`/`w:instrText`, révision supprimée `w:del`/`w:delText`.
"""
import re
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.shared import Cm

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE))

from docie_client import DocIEError, document_payload  # noqa: E402

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ── Construction des DOCX ────────────────────────────────────────────────────

def _python_docx(construire):
    def ecrire(chemin):
        doc = Document()
        construire(doc)
        doc.save(chemin)
    return ecrire


def _xml_a_la_main(corps):
    """Zip minimal : seul `word/document.xml`, seule partie que lit le client."""
    xml = (
        f'<w:document xmlns:w="{W}"'
        ' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
        ' xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        ' xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
        ' xmlns:v="urn:schemas-microsoft-com:vml">'
        f"<w:body>{corps}</w:body></w:document>"
    )

    def ecrire(chemin):
        with ZipFile(chemin, "w") as archive:
            archive.writestr("word/document.xml", xml)
    return ecrire


def _p(*contenu):
    return "<w:p>" + "".join(contenu) + "</w:p>"


def _r(texte):
    return f'<w:r><w:t xml:space="preserve">{texte}</w:t></w:r>'


def _saut_de_ligne(doc):
    doc.add_paragraph("Alice Dupont")
    for lignes in (
        ["alice.dupont@example.com", "06 12 34 56 78", "Lille"],
        ["Data Engineer", "Mars 2022 - Aujourd'hui", "Decathlon, Villeneuve-d'Ascq"],
        ["Developpeuse BI", "Septembre 2018 - Février 2022", "Leroy Merlin, Lezennes"],
    ):
        paragraphe = doc.add_paragraph()
        for i, ligne in enumerate(lignes):
            run = paragraphe.add_run(ligne)
            if i < len(lignes) - 1:
                run.add_break()


def _tabulation(doc):
    paragraphe = doc.add_paragraph()
    paragraphe.add_run("Data Engineer").add_tab()
    paragraphe.add_run("Mars 2022 - Aujourd'hui")
    paragraphe = doc.add_paragraph()
    # Taquet de tabulation : un `w:tab` dans `w:pPr/w:tabs`, qui n'est PAS
    # un caractère et ne doit rien rendre. Pas sur le premier paragraphe : le
    # `.strip()` final y masquerait une tabulation parasite.
    paragraphe.paragraph_format.tab_stops.add_tab_stop(Cm(12))
    paragraphe.add_run("Master Informatique").add_tab()
    paragraphe.add_run("Université de Lille").add_tab()
    paragraphe.add_run("2018")


COMPETENCES = [("Langages", "Python, SQL"), ("Cloud", "AWS, Azure"), ("Outils", "Git, Docker")]


def _tableau(doc):
    doc.add_paragraph("Compétences")
    table = doc.add_table(rows=len(COMPETENCES), cols=2)
    for rangee, (categorie, elements) in zip(table.rows, COMPETENCES):
        rangee.cells[0].text = categorie
        rangee.cells[1].text = elements


def _mise_en_page(doc):
    """Modèle « barre latérale » : 1 rangée, 2 cellules pleines de paragraphes,
    et un tableau de missions imbriqué dans la cellule de droite."""
    table = doc.add_table(rows=1, cols=2)
    gauche, droite = table.rows[0].cells
    gauche.text = "alice.dupont@example.com"
    gauche.add_paragraph("06 12 34 56 78")
    droite.text = "Data Engineer"
    missions = droite.add_table(rows=2, cols=2)
    for rangee, (dates, employeur) in zip(missions.rows, [
        ("Mars 2022 - Aujourd'hui", "Decathlon"),
        ("Septembre 2018 - Février 2022", "Leroy Merlin"),
    ]):
        rangee.cells[0].text = dates
        rangee.cells[1].text = employeur


_RETOUR_CHARIOT = _p('<w:r><w:t>Anglais</w:t><w:cr/><w:t>C1</w:t></w:r>')


def _zone(paragraphes):
    return "<w:txbxContent>" + "".join(_p(_r(t)) for t in paragraphes) + "</w:txbxContent>"


_ZONE_DE_TEXTE = (
    _p(_r("Alice Dupont"))
    + _p(
        "<w:r><mc:AlternateContent>"
        '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData>'
        "<wps:wsp><wps:txbx>" + _zone(["alice.dupont@example.com", "06 12 34 56 78"])
        + "</wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>"
        "<mc:Fallback><w:pict><v:shape><v:textbox>"
        + _zone(["alice.dupont@example.com", "06 12 34 56 78"])
        + "</v:textbox></v:shape></w:pict></mc:Fallback>"
        "</mc:AlternateContent></w:r>"
    )
    + _p(_r("Data Engineer"))
)

_CHAMP = _p(
    '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> HYPERLINK "mailto:alice.dupont@example.com" </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    + _r("alice.dupont@example.com")
    + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
)

_REVISION = _p(
    _r("Septembre 2018 - ")
    # La suppression emporte aussi une tabulation : elle ne doit pas compter.
    + "<w:del><w:r><w:delText>Aujourd'hui</w:delText><w:tab/></w:r></w:del>"
    + "<w:ins>" + _r("Février 2022") + "</w:ins>"
)


# nom → (construction, valeurs connues, textes qui ne doivent PAS apparaître)
CAS = {
    "saut_de_ligne": (_python_docx(_saut_de_ligne), [
        "Alice Dupont", "alice.dupont@example.com", "06 12 34 56 78", "Lille",
        "Data Engineer", "Mars 2022", "Aujourd'hui", "Decathlon",
        "Developpeuse BI", "Septembre 2018", "Février 2022", "Leroy Merlin", "Lezennes",
    ], []),
    "retour_chariot": (_xml_a_la_main(_RETOUR_CHARIOT), ["Anglais", "C1"], []),
    "tabulation": (_python_docx(_tabulation), [
        "Data Engineer", "Mars 2022", "Aujourd'hui",
        "Master Informatique", "Université de Lille", "2018",
    ], []),
    "tableau": (_python_docx(_tableau), [v for rangee in COMPETENCES for v in rangee], []),
    "mise_en_page": (_python_docx(_mise_en_page), [
        "alice.dupont@example.com", "06 12 34 56 78", "Data Engineer",
        "Mars 2022", "Aujourd'hui", "Decathlon", "Septembre 2018", "Février 2022", "Leroy Merlin",
    ], []),
    "zone_de_texte": (_xml_a_la_main(_ZONE_DE_TEXTE),
                      ["Alice Dupont", "alice.dupont@example.com", "06 12 34 56 78", "Data Engineer"], []),
    "code_de_champ": (_xml_a_la_main(_CHAMP), ["alice.dupont@example.com"], ["HYPERLINK", "mailto:"]),
    "revision_supprimee": (_xml_a_la_main(_REVISION), ["Septembre 2018", "Février 2022"], ["Aujourd'hui"]),
}


def ancien_lecteur(chemin):
    """Le lecteur d'avant, recopié tel quel (témoin)."""
    with ZipFile(chemin) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ns = {"w": W}
    return "\n".join("".join(p.itertext()) for p in root.findall(".//w:p", ns)).strip()


def lecteur_actuel(chemin):
    return document_payload(Path(chemin))["text"]


def delimitee(texte, valeur):
    return re.search(r"(?<![0-9A-Za-zÀ-ÿ])" + re.escape(valeur) + r"(?![0-9A-Za-zÀ-ÿ])", texte) is not None


def mesurer(lecteur, nom):
    """(valeurs délimitées, valeurs connues, occurrences max d'une valeur, fuites)."""
    construire, valeurs, interdits = CAS[nom]
    with tempfile.TemporaryDirectory() as dossier:
        chemin = Path(dossier) / f"{nom}.docx"
        construire(chemin)
        texte = lecteur(chemin)
    return {
        "texte": texte,
        "delimitees": sum(delimitee(texte, v) for v in valeurs),
        "connues": len(valeurs),
        "occurrences_max": max(texte.count(v) for v in valeurs),
        "fuites": [i for i in interdits if i in texte],
        "lignes": lignes_docie(texte),
    }


def lignes_docie(texte):
    """Blocs que DocIE tire du texte sur /v1/extract/text : une ligne non vide
    = un bloc (`splitlines()`, lignes vides écartées). Au-delà de 800 blocs,
    hors profils `document_only`/`nuextract*`, la suite n'atteint pas le modèle."""
    return len([ligne for ligne in texte.splitlines() if ligne.strip()])


# ── Lecteur corrigé ──────────────────────────────────────────────────────────

class TexteDocxTests(unittest.TestCase):

    def test_chaque_valeur_connue_survit_delimitee_une_seule_fois(self):
        for nom in CAS:
            with self.subTest(cas=nom):
                m = mesurer(lecteur_actuel, nom)
                self.assertEqual(m["delimitees"], m["connues"], m["texte"])
                self.assertEqual(m["occurrences_max"], 1, m["texte"])
                self.assertEqual(m["fuites"], [], m["texte"])

    def test_saut_de_ligne_rend_un_saut_de_ligne(self):
        texte = mesurer(lecteur_actuel, "saut_de_ligne")["texte"]
        self.assertEqual(texte.split("\n")[:4], ["Alice Dupont", "alice.dupont@example.com", "06 12 34 56 78", "Lille"])
        self.assertIn("Mars 2022 - Aujourd'hui\nDecathlon, Villeneuve-d'Ascq", texte)

    def test_retour_chariot_rend_un_saut_de_ligne(self):
        self.assertEqual(mesurer(lecteur_actuel, "retour_chariot")["texte"], "Anglais\nC1")

    def test_tabulation_rend_une_tabulation_et_le_taquet_rien(self):
        self.assertEqual(
            mesurer(lecteur_actuel, "tabulation")["texte"],
            "Data Engineer\tMars 2022 - Aujourd'hui\nMaster Informatique\tUniversité de Lille\t2018",
        )

    def test_une_rangee_simple_tient_sur_une_ligne(self):
        self.assertEqual(
            mesurer(lecteur_actuel, "tableau")["texte"],
            "Compétences\nLangages\tPython, SQL\nCloud\tAWS, Azure\nOutils\tGit, Docker",
        )

    def test_une_rangee_de_mise_en_page_est_lue_cellule_par_cellule(self):
        """Aplatir la barre latérale en une ligne détruirait toutes ses
        frontières de paragraphe ; le tableau imbriqué, lui, est simple."""
        self.assertEqual(
            mesurer(lecteur_actuel, "mise_en_page")["texte"],
            "alice.dupont@example.com\n06 12 34 56 78\nData Engineer\n"
            "Mars 2022 - Aujourd'hui\tDecathlon\nSeptembre 2018 - Février 2022\tLeroy Merlin",
        )

    def test_zone_de_texte_une_fois_dans_l_ordre_du_document(self):
        self.assertEqual(
            mesurer(lecteur_actuel, "zone_de_texte")["texte"],
            "Alice Dupont\n\nalice.dupont@example.com\n06 12 34 56 78\n\nData Engineer",
        )

    def test_code_de_champ_exclu_resultat_conserve(self):
        self.assertEqual(mesurer(lecteur_actuel, "code_de_champ")["texte"], "alice.dupont@example.com")

    def test_revision_supprimee_exclue(self):
        self.assertEqual(mesurer(lecteur_actuel, "revision_supprimee")["texte"], "Septembre 2018 - Février 2022")

    def test_un_cv_volumineux_reste_sous_le_plafond_texte_docie(self):
        """DocIE refuse au-delà de `max_text_chars` = 1 000 000 caractères."""
        def gros_cv(doc):
            doc.add_paragraph("Alice Dupont")
            for i in range(40):
                p = doc.add_paragraph()
                p.add_run(f"Consultante data senior {i}").add_break()
                p.add_run("Janvier 2010 - Décembre 2011").add_tab()
                p.add_run("Entreprise cliente, Lille")
                for j in range(12):
                    doc.add_paragraph(f"Réalisation {j} : conception et industrialisation "
                                      "de pipelines de données sur Python, Spark, Airflow et Azure.")
            table = doc.add_table(rows=30, cols=2)
            for rangee in table.rows:
                rangee.cells[0].text, rangee.cells[1].text = "Catégorie", "Python, SQL, Scala, Spark, Kafka"
        with tempfile.TemporaryDirectory() as dossier:
            chemin = Path(dossier) / "gros.docx"
            _python_docx(gros_cv)(chemin)
            texte = lecteur_actuel(chemin)
        self.assertGreater(len(texte), 40_000)
        self.assertLess(len(texte), 1_000_000)


class GardesInchangeesTests(unittest.TestCase):

    def _ecrire(self, contenu):
        dossier = tempfile.TemporaryDirectory()
        self.addCleanup(dossier.cleanup)
        chemin = Path(dossier.name) / "cv.docx"
        contenu(chemin)
        return chemin

    def test_document_sans_texte_lisible(self):
        """Des tabulations et un code de champ seuls ne sont pas du texte."""
        chemin = self._ecrire(_xml_a_la_main(
            _p('<w:r><w:tab/><w:br/></w:r>')
            + _p('<w:r><w:instrText> PAGE </w:instrText></w:r>')))
        with self.assertRaises(DocIEError) as ctx:
            document_payload(chemin)
        self.assertEqual(str(ctx.exception),
                         "Document Word sans texte lisible : exportez-le en PDF pour l'OCR de la plateforme d'inférence interne.")

    def test_document_word_invalide(self):
        for contenu in (b"pas un zip", None):
            with self.subTest(contenu=contenu):
                def ecrire(chemin, contenu=contenu):
                    if contenu is None:
                        with ZipFile(chemin, "w") as archive:
                            archive.writestr("word/document.xml", "<w:document")
                    else:
                        chemin.write_bytes(contenu)
                with self.assertRaises(DocIEError) as ctx:
                    document_payload(self._ecrire(ecrire))
                self.assertEqual(str(ctx.exception), "Document Word invalide.")

    def test_document_trop_volumineux_apres_decompression(self):
        def ecrire(chemin):
            with ZipFile(chemin, "w", compression=8) as archive:
                archive.writestr("word/document.xml", " " * (40 * 1024 * 1024 + 1))
        with self.assertRaises(DocIEError) as ctx:
            document_payload(self._ecrire(ecrire))
        self.assertEqual(str(ctx.exception), "Document Word trop volumineux après décompression.")


# ── Témoin : l'ancien lecteur perdait, sur les mêmes entrées ────────────────

class TemoinAncienLecteurTests(unittest.TestCase):
    """Si l'un de ces tests échoue, c'est que le témoin ne reproduit plus la
    perte : la preuve que le lecteur corrigé la répare n'en est plus une."""

    def test_saut_de_ligne_collait_courriel_telephone_ville(self):
        m = mesurer(ancien_lecteur, "saut_de_ligne")
        self.assertIn("alice.dupont@example.com06 12 34 56 78Lille", m["texte"])
        self.assertIn("Aujourd'huiDecathlon", m["texte"])
        self.assertLess(m["delimitees"], m["connues"])

    def test_retour_chariot_collait(self):
        self.assertEqual(mesurer(ancien_lecteur, "retour_chariot")["texte"], "AnglaisC1")

    def test_tabulation_collait(self):
        m = mesurer(ancien_lecteur, "tabulation")
        self.assertIn("Data EngineerMars 2022", m["texte"])
        self.assertIn("Université de Lille2018", m["texte"])

    def test_tableau_separait_categorie_et_elements(self):
        texte = mesurer(ancien_lecteur, "tableau")["texte"]
        self.assertIn("Langages\nPython, SQL", texte)
        self.assertNotIn("Langages\tPython, SQL", texte)

    def test_zone_de_texte_comptee_quatre_fois(self):
        m = mesurer(ancien_lecteur, "zone_de_texte")
        self.assertEqual(m["texte"].count("06 12 34 56 78"), 4)

    def test_code_de_champ_et_revision_supprimee_fuyaient(self):
        self.assertEqual(mesurer(ancien_lecteur, "code_de_champ")["fuites"], ["HYPERLINK", "mailto:"])
        m = mesurer(ancien_lecteur, "revision_supprimee")
        self.assertIn("Aujourd'huiFévrier 2022", m["texte"])
        self.assertEqual(m["fuites"], ["Aujourd'hui"])


if __name__ == "__main__":
    unittest.main()
